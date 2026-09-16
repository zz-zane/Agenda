"""Local-only UI preview and disk persistence. Run: python local_server.py."""
from datetime import date, datetime
from contextlib import contextmanager
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
import argparse
import hashlib
from itertools import islice
import json
import mimetypes
import re
import sqlite3
import uuid
import zipfile
import sys
from urllib.parse import unquote
import local_ai
import local_profile

from PIL import Image, ImageOps, UnidentifiedImageError
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
DATA = ROOT / "data"
LIMIT = 12 * 1024 * 1024


def today():
    return date.today().isoformat()


@contextmanager
def database():
    db = sqlite3.connect(DATA / "calendar.db", timeout=10)
    db.row_factory = sqlite3.Row
    try:
        with db:
            yield db
    finally:
        db.close()


def initialize():
    (DATA / "photos").mkdir(parents=True, exist_ok=True)
    with database() as db:
        db.executescript("""
        CREATE TABLE IF NOT EXISTS tasks(id TEXT PRIMARY KEY, date TEXT NOT NULL,
          title TEXT NOT NULL, start TEXT NOT NULL, end TEXT NOT NULL,
          note TEXT NOT NULL DEFAULT '', origin TEXT UNIQUE);
        CREATE TABLE IF NOT EXISTS photos(id TEXT PRIMARY KEY, date TEXT NOT NULL, filename TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS checkins(date TEXT PRIMARY KEY, at TEXT NOT NULL);
        CREATE TABLE IF NOT EXISTS visits(date TEXT PRIMARY KEY);
        """)
        local_ai.initialize(db)


def valid_date(value):
    if not isinstance(value, str) or not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        raise ValueError("日期格式不正确")
    date.fromisoformat(value)
    return value


def minutes(value):
    if not isinstance(value, str) or not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d|24:00", value):
        raise ValueError("时间格式不正确")
    h, m = map(int, value.split(":"))
    return h * 60 + m


def task_fields(body):
    day = valid_date(body.get("date"))
    title = body.get("title", "")
    note = body.get("note", "")
    if not isinstance(title, str) or not 1 <= len(title.strip()) <= 100:
        raise ValueError("请填写 1–100 字的任务名称")
    if not isinstance(note, str) or len(note) > 1000:
        raise ValueError("备注不能超过 1000 字")
    start, end = body.get("start"), body.get("end")
    if not 0 <= minutes(start) < minutes(end) <= 1440:
        raise ValueError("结束时间必须晚于开始时间，最晚为 24:00")
    return day, title.strip(), start, end, note


def parse_xlsx(raw):
    """Use this project's actual 日程明细 sheet, or its dated weekly grids."""
    with zipfile.ZipFile(BytesIO(raw)) as archive:
        if len(archive.infolist()) > 300 or sum(i.file_size for i in archive.infolist()) > 64 * 1024 * 1024:
            raise ValueError("课表文件解压后过大")
    workbook = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    tasks = []
    try:
        # The supplied workbook includes a normalized detail sheet with combined periods.
        sheets = [workbook["日程明细"]] if "日程明细" in workbook.sheetnames else workbook.worksheets
        for sheet in sheets:
            if (sheet.max_row or 0) > 10000 or (sheet.max_column or 0) > 100:
                raise ValueError("课表行数或列数过多")
            rows = list(islice(sheet.values, 10001))
            if len(rows) > 10000 or any(len(r) > 100 for r in rows):
                raise ValueError("课表行数或列数过多")
            header = next((i for i, r in enumerate(rows[:20]) if "日期" in r and ("课程" in r or "课程名称" in r)), None)
            if header is not None:
                columns = {str(v).strip(): i for i, v in enumerate(rows[header]) if v is not None}
                def value(row, *names):
                    return next((row[columns[n]] for n in names if n in columns and columns[n] < len(row)), None)
                for row in rows[header + 1:]:
                    d, title = value(row, "日期"), value(row, "课程", "课程名称", "标题")
                    if d is None and title is None:
                        continue
                    if isinstance(d, (date, datetime)):
                        d = d.strftime("%Y-%m-%d")
                    span = str(value(row, "时间") or "")
                    note = " · ".join(str(v) for v in [value(row, "地点"), value(row, "备注")] if v)
                    # A course can have separate morning/afternoon spans; keep lunch free.
                    spans = re.split(r"[；;]", span) if span else [""]
                    for part in spans:
                        pair = re.fullmatch(r"\s*(\d{2}:\d{2})\s*[-–—~至]\s*(\d{2}:\d{2})\s*", part)
                        start = pair[1] if pair else str(value(row, "开始时间") or "")[:5]
                        end = pair[2] if pair else str(value(row, "结束时间") or "")[:5]
                        tasks.append(task_fields(dict(date=str(d), title=str(title or ""), start=start, end=end, note=note)))
            else:
                # Weekly sheets carry actual ISO dates in day headers and time spans in column C.
                index = next((i for i, r in enumerate(rows[:10]) if len(r) > 3 and str(r[2]) == "具体时间"), None)
                if index is None:
                    continue
                days = {i: re.search(r"\d{4}-\d{2}-\d{2}", str(v))[0] for i, v in enumerate(rows[index]) if re.search(r"\d{4}-\d{2}-\d{2}", str(v))}
                for row in rows[index + 1:]:
                    pair = re.fullmatch(r"(\d{2}:\d{2})-(\d{2}:\d{2})", str(row[2] if len(row) > 2 else ""))
                    if not pair:
                        continue
                    for col, day in days.items():
                        if col < len(row) and row[col]:
                            lines = str(row[col]).splitlines()
                            tasks.append(task_fields(dict(date=day, title=lines[0], start=pair[1], end=pair[2], note=" · ".join(lines[1:]))))
        if not tasks:
            raise ValueError("未找到带具体日期和起止时间的课表，请选择项目中的学期 XLSX")
        if len(tasks) > 10000:
            raise ValueError("课表超过 10000 条")
        return list(dict.fromkeys(tasks))
    finally:
        workbook.close()


class Handler(SimpleHTTPRequestHandler):
    def json(self, value, status=200):
        raw = json.dumps(value, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(raw)))
        self.end_headers()
        self.wfile.write(raw)

    def local_request(self):
        host = self.headers.get("Host", "")
        allowed = {f"127.0.0.1:{self.server.server_port}", f"localhost:{self.server.server_port}"}
        origin = self.headers.get("Origin")
        return host in allowed and (not origin or origin == "http://" + host)

    def do_GET(self):
        if not self.local_request():
            return self.json({"error": "仅允许本机访问"}, 403)
        path = self.path.split("?", 1)[0]
        if path == "/api/state":
            with database() as db:
                result = {key: [dict(r) for r in db.execute("SELECT * FROM " + key)] for key in ("tasks", "photos", "checkins", "visits")}
                result['tasks'] = [t for t in local_ai.context(db, sys.modules[__name__])['tasks'] if not t.get('superseded')]
            return self.json({**result, "today": today()})
        if path == '/api/ai/state':
            return self.json(local_ai.state(sys.modules[__name__]))
        if path == '/api/profile':
            with database() as db:return self.json(local_profile.state(db))
        if path.startswith("/photos/"):
            name = path.removeprefix("/photos/")
            if not re.fullmatch(r"[a-f0-9]{32}\.jpg", name):
                return self.json({"error": "照片不存在"}, 404)
            file = DATA / "photos" / name
        else:
            names = {"/": "index.html", "/index.html": "index.html", "/app.js": "app.js", "/ai.js": "ai.js", "/i18n.js": "i18n.js", "/calendar.js": "calendar.js", "/styles.css": "styles.css", "/favicon.svg": "favicon.svg"}
            if path not in names:
                return self.json({"error": "未找到页面"}, 404)
            file = ROOT / "frontend" / names[path]
        if not file.is_file():
            return self.json({"error": "文件不存在"}, 404)
        raw = file.read_bytes()
        self.send_response(200)
        self.send_header("Content-Type", mimetypes.guess_type(file.name)[0] or "application/octet-stream")
        self.send_header("Content-Length", str(len(raw)))
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Content-Type-Options", "nosniff")
        self.end_headers()
        self.wfile.write(raw)

    def do_POST(self):
        if not self.local_request():
            return self.json({"error": "仅允许本机页面保存"}, 403)
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if not 0 < length <= LIMIT:
                return self.json({"error": "文件不能为空且不能超过 12MB"}, 413)
            raw = self.rfile.read(length)
            route = self.path
            if route in ("/api/import", "/api/ai/attachment"):
                tasks = parse_xlsx(raw)
                name = Path(unquote(self.headers.get('X-Filename','课表.xlsx'))).name[:200]
                with database() as db:
                    if route == '/api/ai/attachment':
                        attachment=uuid.uuid4().hex
                        db.execute('INSERT INTO attachments VALUES (?,?,?)',(attachment,name,json.dumps(tasks,ensure_ascii=False)))
                        result={'id':attachment,'name':name,'count':len(tasks)}
                        local_profile.archive(db,name,raw,'attachment',attachment)
                    else:
                        result=local_ai.import_tasks(db,tasks,name)
                        local_profile.archive(db,name,raw,'import',result['batch_id'])
                return self.json(result)
            if route.startswith("/api/photos/"):
                day = valid_date(route.rsplit("/", 1)[-1])
                if day != today():
                    return self.json({"error": "只有当天可以上传照片，历史日期只读"}, 403)
                with Image.open(BytesIO(raw)) as source:
                    if source.format not in {"JPEG", "PNG", "WEBP"} or source.width * source.height > 24_000_000:
                        raise ValueError("请选择 2400 万像素以内的 JPG、PNG 或 WebP 照片")
                    picture = ImageOps.exif_transpose(source).convert("RGB")
                    picture.thumbnail((2400, 2400))
                    name = uuid.uuid4().hex + ".jpg"
                    target = DATA / "photos" / name
                    # Check again after image decoding, including across midnight.
                    if day != today():
                        raise ValueError("日期已变化，请重新打开今天")
                    picture.save(target, "JPEG", quality=90)
                try:
                    with database() as db:
                        db.execute("BEGIN IMMEDIATE")
                        if day != today():
                            raise ValueError("日期已变化，请重新打开今天")
                        db.execute("INSERT INTO photos VALUES (?,?,?)", (name[:-4], day, name))
                except Exception:
                    target.unlink(missing_ok=True)
                    raise
                return self.json({"filename": name})
            body = json.loads(raw)
            if not isinstance(body, dict):
                raise ValueError("请求格式错误")
            app = sys.modules[__name__]
            if route == '/api/profile/generate':
                return self.json(local_profile.generate(app))
            if route == '/api/ai/config/select':
                return self.json(local_ai.select_config(app,body))
            if route == '/api/ai/config':
                return self.json(local_ai.save_config(app,body))
            if route == '/api/ai/chat':
                return self.json(local_ai.chat(app,body))
            if route == '/api/ai/apply':
                return self.json(local_ai.apply(app,body.get('id')))
            with database() as db:
                db.execute("BEGIN IMMEDIATE")
                if route == "/api/open":
                    db.execute("INSERT OR IGNORE INTO visits VALUES (?)", (today(),))
                elif route == "/api/checkin":
                    day = valid_date(body.get("date"))
                    if day != today():
                        return self.json({"error": "只能完成当天打卡"}, 403)
                    if not db.execute("SELECT 1 FROM photos WHERE date=?", (day,)).fetchone():
                        raise ValueError("请先上传当天照片，再完成打卡")
                    db.execute("INSERT OR IGNORE INTO checkins VALUES (?,?)", (day, datetime.now().isoformat()))
                elif route == "/api/tasks":
                    fields = task_fields(body)
                    if fields[0] < today():
                        return self.json({"error": "历史日程只可查看"}, 403)
                    task_id = body.get("id")
                    if task_id:
                        old = db.execute("SELECT * FROM tasks WHERE id=?", (task_id,)).fetchone()
                        if not old:
                            raise ValueError("任务不存在，请刷新")
                        if old["date"] < today():
                            return self.json({"error": "不能修改历史日程"}, 403)
                        meta=db.execute('SELECT * FROM task_meta WHERE task_id=?',(task_id,)).fetchone()
                        if meta and (meta['done'] or meta['goal_id'] or meta['superseded']):
                            raise ValueError('已完成任务只读，目标任务请在 AI 中重排以保留大纲和截止日')
                        db.execute("UPDATE tasks SET date=?,title=?,start=?,end=?,note=? WHERE id=?", (*fields, task_id))
                    else:
                        db.execute("INSERT INTO tasks VALUES (?,?,?,?,?,?,NULL)", (uuid.uuid4().hex, *fields))
                elif route == "/api/tasks/delete":
                    old = db.execute("SELECT * FROM tasks WHERE id=?", (body.get("id"),)).fetchone()
                    if not old or old["date"] < today():
                        return self.json({"error": "历史日程不能删除或任务已不存在"}, 403)
                    meta=db.execute('SELECT * FROM task_meta WHERE task_id=?',(old['id'],)).fetchone()
                    if meta and (meta['done'] or meta['goal_id']):
                        raise ValueError('已完成或目标任务不能单独删除')
                    db.execute("DELETE FROM tasks WHERE id=?", (old["id"],))
                    db.execute('DELETE FROM task_meta WHERE task_id=?',(old['id'],))
                elif route == '/api/tasks/done':
                    task=db.execute('SELECT * FROM tasks WHERE id=?',(body.get('id'),)).fetchone()
                    if not task or task['date'] != today(): raise ValueError('只能完成今天的任务')
                    meta=db.execute('SELECT * FROM task_meta WHERE task_id=?',(task['id'],)).fetchone()
                    if task['origin'] or (meta and meta['superseded']): raise ValueError('课表或已替换任务不能标记完成')
                    if not db.execute('SELECT 1 FROM photos WHERE date=?',(today(),)).fetchone(): raise ValueError('请先在月历上传今天的照片')
                    db.execute('INSERT INTO task_meta(task_id,done) VALUES (?,1) ON CONFLICT(task_id) DO UPDATE SET done=1',(task['id'],))
                else:
                    return self.json({"error": "接口不存在"}, 404)
            return self.json({"ok": True})
        except (ValueError, TypeError, KeyError, zipfile.BadZipFile, UnidentifiedImageError, Image.DecompressionBombError) as exc:
            return self.json({"error": str(exc) or "文件无法识别"}, 422)
        except Exception:
            return self.json({"error": "保存失败，请检查文件或本地磁盘后重试"}, 500)


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--data", type=Path, default=DATA)
    args = parser.parse_args()
    DATA = args.data.resolve()
    initialize()
    print(f"Agenda local: http://127.0.0.1:{args.port} | Photos: {DATA / 'photos'}", flush=True)
    ThreadingHTTPServer(("127.0.0.1", args.port), Handler).serve_forever()
